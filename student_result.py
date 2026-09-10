import openpyxl
import os


FILE_NAME = "student_results.xlsx"


# Create Excel file if it does not exist
def create_excel_file():
    if not os.path.exists(FILE_NAME):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "Student Results"

        headers = [
            "Roll No", "Name", "Class",
            "Subject 1", "Subject 2", "Subject 3",
            "Subject 4", "Subject 5",
            "Total", "Percentage", "Grade", "Status"
        ]

        sheet.append(headers)
        workbook.save(FILE_NAME)


# Calculate result
def calculate_result(marks):
    total = sum(marks)
    percentage = total / 5

    # Pass only if every subject has at least 35 marks
    if all(mark >= 35 for mark in marks):
        status = "PASS"

        if percentage >= 90:
            grade = "A+"
        elif percentage >= 80:
            grade = "A"
        elif percentage >= 70:
            grade = "B"
        elif percentage >= 60:
            grade = "C"
        elif percentage >= 50:
            grade = "D"
        elif percentage >= 35:
            grade = "E"
    else:
        grade = "F"
        status = "FAIL"

    return total, percentage, grade, status


# Add student
def add_student():
    print("\n--------------------------------")
    print("       ADD STUDENT RESULT")
    print("--------------------------------")

    roll_no = input("Enter Roll No: ")
    name = input("Enter Student Name: ")
    student_class = input("Enter Class/Course: ")

    marks = []

    for i in range(1, 6):
        while True:
            try:
                mark = float(input(f"Enter marks of Subject {i}: "))

                if 0 <= mark <= 100:
                    marks.append(mark)
                    break
                else:
                    print("Marks must be between 0 and 100.")

            except ValueError:
                print("Please enter a valid number.")

    total, percentage, grade, status = calculate_result(marks)

    workbook = openpyxl.load_workbook(FILE_NAME)
    sheet = workbook.active

    # Check duplicate roll number
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if str(row[0]) == roll_no:
            print("\nRoll No already exists!")
            workbook.close()
            return

    sheet.append([
        roll_no,
        name,
        student_class,
        marks[0],
        marks[1],
        marks[2],
        marks[3],
        marks[4],
        total,
        percentage,
        grade,
        status
    ])

    workbook.save(FILE_NAME)
    workbook.close()

    print("\nStudent result added successfully!")
    print("--------------------------------")
    print("Total      :", total)
    print("Percentage :", f"{percentage:.2f}%")
    print("Grade      :", grade)
    print("Status     :", status)
    print("--------------------------------")


# Get one student's result
def get_result():
    print("\n--------------------------------")
    print("        GET STUDENT RESULT")
    print("--------------------------------")

    roll_no = input("Enter Roll No: ")

    workbook = openpyxl.load_workbook(FILE_NAME)
    sheet = workbook.active

    found = False

    for row in sheet.iter_rows(min_row=2, values_only=True):

        if str(row[0]) == roll_no:
            found = True

            print("\n--------------------------------")
            print("          STUDENT RESULT")
            print("--------------------------------")
            print("Roll No     :", row[0])
            print("Name        :", row[1])
            print("Class       :", row[2])
            print("Total       :", row[8])
            print("Percentage  :", f"{row[9]:.2f}%")
            print("Grade       :", row[10])
            print("Status      :", row[11])
            print("--------------------------------")

            break

    workbook.close()

    if not found:
        print("\nStudent with this Roll No was not found.")


# Show all student data
def show_all_data():
    print("\n--------------------------------")
    print("       ALL STUDENT DATA")
    print("--------------------------------")

    workbook = openpyxl.load_workbook(FILE_NAME)
    sheet = workbook.active

    if sheet.max_row <= 1:
        print("No student records found.")
        workbook.close()
        return

    print(
        f"{'Roll No':<10}"
        f"{'Name':<15}"
        f"{'Class':<12}"
        f"{'Total':<10}"
        f"{'Percentage':<13}"
        f"{'Grade':<8}"
        f"{'Status':<8}"
    )

    print("-" * 76)

    for row in sheet.iter_rows(min_row=2, values_only=True):

        print(
            f"{str(row[0]):<10}"
            f"{str(row[1]):<15}"
            f"{str(row[2]):<12}"
            f"{str(row[8]):<10}"
            f"{row[9]:<13.2f}"
            f"{str(row[10]):<8}"
            f"{str(row[11]):<8}"
        )

    workbook.close()


# Main menu
def menu():

    create_excel_file()

    while True:

        print("\n========================================")
        print("       STUDENT RESULT MANAGEMENT")
        print("========================================")
        print("1. Add Student Result")
        print("2. Get Student Result")
        print("3. Show All Student Data")
        print("4. Exit")
        print("========================================")

        choice = input("Enter your choice: ")

        if choice == "1":
            add_student()

        elif choice == "2":
            get_result()

        elif choice == "3":
            show_all_data()

        elif choice == "4":
            print("\nThank you for using Student Result Management System!")
            break

        else:
            print("\nInvalid choice! Please enter 1 to 4.")


# Start program
menu()
